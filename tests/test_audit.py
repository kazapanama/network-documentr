"""
Tests for audit.py — what is left unconfigured on Cisco / FortiGate / FortiSwitch.
"""
import textwrap

from audit import (
    run_audit, report_text, summary_text, short_summary, compress_ports,
    range_cmds, cisco_blocks, forti_tree, fedits, plural,
)
from main import parse_config, read_file_safe


def audit(text):
    return run_audit(text, parse_config(text))


# ─────────────────────────────────────────────────────────────────────────────
#  Sample configs
# ─────────────────────────────────────────────────────────────────────────────

CISCO_MESSY = textwrap.dedent("""\
    Building configuration...
    !
    version 15.2
    no service password-encryption
    !
    hostname Switch
    !
    enable password cisco123
    username admin password 0 admin
    !
    vlan 10
     name USERS
    !
    vlan 77
     name OLD
    !
    interface GigabitEthernet0/1
     description PC
     switchport mode access
     switchport access vlan 10
     switchport port-security maximum 2
    !
    interface GigabitEthernet0/2
     switchport mode access
    !
    interface GigabitEthernet0/3
    !
    interface GigabitEthernet0/4
    !
    interface GigabitEthernet0/5
     description Printer
     switchport access vlan 30
     switchport mode access
     spanning-tree portfast
     spanning-tree bpduguard enable
    !
    interface GigabitEthernet0/6
     description to-core
     switchport mode dynamic desirable
    !
    interface GigabitEthernet0/24
     description Uplink
     switchport mode trunk
    !
    interface Vlan1
     ip address 192.168.1.10 255.255.255.0
    !
    ip http server
    snmp-server community public RO
    snmp-server community secretRW RW
    banner motd ^C
    Authorized access only
    hostname FAKE
    ^C
    !
    line con 0
     exec-timeout 0 0
    line vty 0 4
     password cisco
     login
     exec-timeout 0 0
    line vty 5 15
     transport input telnet ssh
     login local
    !
    end
""")

CISCO_CLEAN = textwrap.dedent("""\
    Building configuration...
    !
    service timestamps log datetime msec localtime show-timezone
    service password-encryption
    !
    hostname SW-GOOD
    !
    enable secret 9 $9$abcdefghijk
    username netadmin privilege 15 secret 9 $9$abcdefghijk
    clock timezone EET 2 0
    no ip domain lookup
    ip dhcp snooping
    spanning-tree mode rapid-pvst
    spanning-tree portfast bpduguard default
    errdisable recovery cause bpduguard
    !
    vlan 10
     name USERS
    !
    vlan 99
     name MGMT
    !
    interface GigabitEthernet1/0/1
     description PC
     switchport access vlan 10
     switchport mode access
     spanning-tree portfast
    !
    interface GigabitEthernet1/0/2
     description unused
     switchport access vlan 10
     switchport mode access
     shutdown
    !
    interface GigabitEthernet1/0/48
     description Uplink
     switchport trunk native vlan 999
     switchport trunk allowed vlan 10,99
     switchport mode trunk
     switchport nonegotiate
    !
    interface Vlan99
     description MGMT
     ip address 10.0.99.2 255.255.255.0
    !
    ip default-gateway 10.0.99.1
    ip ssh version 2
    no ip http server
    no ip http secure-server
    ntp server 10.0.0.1
    logging host 10.0.0.5
    snmp-server community Xk9qPz RO 10
    banner login ^C
    Authorized access only
    ^C
    !
    line con 0
     exec-timeout 10 0
    line vty 0 15
     access-class 10 in
     exec-timeout 10 0
     login local
     transport input ssh
    !
    end
""")

FGT_MESSY = textwrap.dedent("""\
    #config-version=FGT60F-7.2.5-FW-build1517-230606:opmode=0:vdom=0:user=admin
    config system global
        set admintimeout 480
        set hostname "FGT60F"
        set strong-crypto disable
        set admin-https-ssl-versions tlsv1-1 tlsv1-2
    end
    config system interface
        edit "wan1"
            set ip 203.0.113.2 255.255.255.252
            set allowaccess ping https ssh http
            set role wan
        next
        edit "lan"
            set ip 192.168.1.1 255.255.255.0
            set allowaccess ping https ssh telnet
            set role lan
        next
        edit "dmz"
            set ip 10.10.10.1 255.255.255.0
            set alias "Servers"
        next
    end
    config system admin
        edit "admin"
            set accprofile "super_admin"
            set trusthost1 0.0.0.0 0.0.0.0
        next
        edit "ivan"
            set accprofile "super_admin"
            set trusthost1 10.0.0.0 255.255.255.0
            set two-factor fortitoken
        next
    end
    config system ntp
        set ntpsync disable
    end
    config system snmp community
        edit 1
            set name "public"
            config hosts
                edit 1
                    set ip 0.0.0.0 0.0.0.0
                next
            end
        next
    end
    config system password-policy
        set status disable
    end
    config log syslogd setting
        set status disable
    end
    config log fortianalyzer setting
        set status disable
    end
    config firewall policy
        edit 1
            set name "ANY-ANY"
            set srcintf "lan"
            set dstintf "wan1"
            set srcaddr "all"
            set dstaddr "all"
            set action accept
            set schedule "always"
            set service "ALL"
            set logtraffic disable
            set nat enable
        next
        edit 2
            set srcintf "lan"
            set dstintf "dmz"
            set srcaddr "all"
            set dstaddr "all"
            set action accept
            set schedule "always"
            set service "ALL"
            set status disable
        next
    end
""")

FGT_CLEAN = textwrap.dedent("""\
    #config-version=FGT60F-7.2.5-FW-build1517-230606:opmode=0:vdom=0:user=admin
    config system global
        set admintimeout 10
        set hostname "FGT-HQ"
    end
    config system interface
        edit "wan1"
            set ip 203.0.113.2 255.255.255.252
            set allowaccess ping
            set role wan
            set alias "ISP"
        next
        edit "lan"
            set ip 192.168.1.1 255.255.255.0
            set allowaccess ping https ssh
            set alias "Office"
        next
    end
    config system admin
        edit "o.koval"
            set trusthost1 192.168.1.0 255.255.255.0
            set two-factor fortitoken
        next
    end
    config system ntp
        set ntpsync enable
        set type fortiguard
    end
    config system password-policy
        set status enable
    end
    config log syslogd setting
        set status enable
        set server "10.0.0.5"
    end
""")

FSW = textwrap.dedent("""\
    #config-version=S124EP-7.2.3-FW-build0450-230317:opmode=0:vdom=0:user=admin
    config system global
        set hostname "FSW-FLOOR1"
    end
    config switch physical-port
        edit "port1"
            set description "PC"
        next
        edit "port2"
        next
        edit "port3"
            set status down
        next
    end
    config switch interface
        edit "port1"
            set native-vlan 10
        next
        edit "port2"
            set allowed-vlans-all enable
            set stp-state disabled
        next
    end
""")


# ─────────────────────────────────────────────────────────────────────────────
#  Helpers
# ─────────────────────────────────────────────────────────────────────────────

class TestHelpers:
    def test_compress_ports(self):
        names = ["GigabitEthernet1/0/1", "GigabitEthernet1/0/2", "GigabitEthernet1/0/3",
                 "GigabitEthernet1/0/9", "Port-channel1"]
        assert compress_ports(names) == "Gi1/0/1-3, Gi1/0/9, Po1"

    def test_compress_ports_keeps_stack_members_apart(self):
        assert compress_ports(["Gi1/0/48", "Gi2/0/1"]) == "Gi1/0/48, Gi2/0/1"

    def test_range_cmds_single(self):
        assert range_cmds(["GigabitEthernet0/1"], ["shutdown"]) == "interface Gi0/1\n shutdown"

    def test_range_cmds_range(self):
        out = range_cmds(["Gi0/1", "Gi0/2", "Gi0/5"], ["spanning-tree portfast"])
        assert out == "interface range Gi0/1 - 2 , Gi0/5\n spanning-tree portfast"

    def test_range_cmds_chunks_of_five(self):
        names = [f"Gi0/{i}" for i in range(1, 24, 2)]  # 12 одиночних портів
        out = range_cmds(names, ["shutdown"])
        assert out.count("interface range") == 3

    def test_plural(self):
        assert plural(1, "a", "b", "c") == "a"
        assert plural(3, "a", "b", "c") == "b"
        assert plural(11, "a", "b", "c") == "c"
        assert plural(22, "a", "b", "c") == "b"
        assert plural(25, "a", "b", "c") == "c"

    def test_cisco_blocks_skips_multiline_banner(self):
        heads = [h for h, _ in cisco_blocks(CISCO_MESSY)]
        assert "hostname FAKE" not in heads
        assert "line con 0" in heads

    def test_cisco_blocks_stops_at_end(self):
        text = "hostname A\ninterface Gi0/1\n description x\nend\nshow inventory\nPID: X"
        heads = [h for h, _ in cisco_blocks(text)]
        assert heads == ["hostname A", "interface Gi0/1"]

    def test_forti_tree_nested_and_quotes(self):
        root = forti_tree(FGT_MESSY)
        comms = fedits(root, "system snmp community")
        assert comms[0]["attrs"]["name"] == ["public"]
        hosts = [c for c in comms[0]["children"] if c["name"] == "hosts"]
        assert hosts and hosts[0]["children"][0]["attrs"]["ip"] == ["0.0.0.0", "0.0.0.0"]

    def test_forti_tree_multiline_value(self):
        text = ('config vpn certificate local\n    edit "c"\n        set private-key "-----BEGIN\nAAA\n-----END"\n'
                '        set comments "after"\n    next\nend\nconfig system global\n    set hostname "X"\nend\n')
        root = forti_tree(text)
        assert fedits(root, "vpn certificate local")[0]["attrs"]["comments"] == ["after"]

    def test_forti_tree_multi_vdom(self):
        text = textwrap.dedent("""\
            config vdom
            edit root
            config firewall policy
                edit 1
                    set action accept
                next
            end
            next
            edit dmz
            config firewall policy
                edit 7
                    set action deny
                next
            end
            next
            end
        """)
        assert [e["name"] for e in fedits(forti_tree(text), "firewall policy")] == ["1", "7"]


# ─────────────────────────────────────────────────────────────────────────────
#  Cisco
# ─────────────────────────────────────────────────────────────────────────────

class TestCiscoMessy:
    rep = audit(CISCO_MESSY)

    def test_vendor(self):
        assert self.rep.vendor == "Cisco IOS"

    def test_expected_codes(self):
        expected = {
            "enable-password", "user-password", "password-encryption", "vty-telnet",
            "vty-transport-default", "vty-no-acl", "vty-no-timeout", "vty-shared-password", "ssh-v1",
            "con-no-timeout", "http-server", "snmp-default-community", "snmp-rw", "snmp-no-acl",
            "no-ntp", "no-syslog", "no-timestamps", "no-timezone", "default-hostname", "domain-lookup",
            "stp-pvst", "port-empty", "port-dtp", "access-vlan1", "access-no-bpduguard",
            "access-no-portfast", "trunk-all-vlans", "trunk-native1", "trunk-dtp", "portsec-not-enabled",
            "no-description", "vlan-undefined", "vlan-unused", "svi-vlan1", "no-dhcp-snooping",
            "no-errdisable-recovery", "no-default-gateway",
        }
        assert expected <= self.rep.codes(), expected - self.rep.codes()

    def test_banner_present_not_flagged(self):
        assert "no-banner" not in self.rep.codes()

    def test_sorted_by_severity(self):
        order = {"high": 0, "medium": 1, "low": 2}
        sevs = [order[f.sev] for f in self.rep.findings]
        assert sevs == sorted(sevs)

    def test_empty_ports(self):
        assert self.rep.get("port-empty").where == ["Gi0/3-4"]
        assert "interface range Gi0/3 - 4" in self.rep.get("port-empty").fix

    def test_dtp_includes_dynamic_mode(self):
        assert self.rep.get("port-dtp").where == ["Gi0/6"]

    def test_access_vlan1(self):
        assert self.rep.get("access-vlan1").where == ["Gi0/2"]

    def test_bpduguard_only_missing_ports(self):
        assert self.rep.get("access-no-bpduguard").where == ["Gi0/1-2"]

    def test_portsec_options_without_enable(self):
        assert self.rep.get("portsec-not-enabled").where == ["Gi0/1"]

    def test_vlan_undefined(self):
        assert self.rep.get("vlan-undefined").where == ["VLAN 30: Gi0/5"]

    def test_vlan_unused(self):
        assert self.rep.get("vlan-unused").where == ["77 (OLD)"]

    def test_vty_findings_name_lines(self):
        assert self.rep.get("vty-telnet").where == ["line vty 5 15"]
        assert self.rep.get("vty-transport-default").where == ["line vty 0 4"]
        assert self.rep.get("vty-no-acl").where == ["line vty 0 4", "line vty 5 15"]

    def test_snmp(self):
        assert self.rep.get("snmp-default-community").where == ["public"]
        assert self.rep.get("snmp-rw").where == ["secretRW"]
        assert set(self.rep.get("snmp-no-acl").where) == {"public", "secretRW"}

    def test_user_password(self):
        assert self.rep.get("user-password").where == ["admin"]

    def test_vty_login_local_has_users(self):
        assert "vty-login-local-no-users" not in self.rep.codes()

    def test_report_text(self):
        txt = report_text(self.rep, "Switch")
        assert txt.startswith("АУДИТ КОНФІГУРАЦІЇ — Switch (Cisco IOS)")
        assert "[КРИТИЧНО] Пароль привілейованого режиму" in txt
        assert "      line vty 0 15" in txt

    def test_summaries(self):
        c = self.rep.counts()
        assert summary_text(self.rep).startswith(f"{c['high']} критичн")
        assert short_summary(self.rep).startswith("Аудит: ")


class TestCiscoClean:
    rep = audit(CISCO_CLEAN)

    def test_no_findings(self):
        assert self.rep.findings == [], [f.code for f in self.rep.findings]
        assert summary_text(self.rep) == "зауважень немає"
        assert short_summary(self.rep) == "Аудит: OK"
        assert self.rep.skipped == []

    def test_global_portfast_default_satisfies_bpduguard(self):
        text = CISCO_CLEAN.replace(" spanning-tree portfast\n", "")
        text = text.replace("spanning-tree portfast bpduguard default", "spanning-tree portfast default\n"
                                                                          "spanning-tree portfast bpduguard default")
        rep = audit(text)
        assert "access-no-bpduguard" not in rep.codes()
        assert "access-no-portfast" not in rep.codes()


class TestCiscoEdgeCases:
    def test_fragment_skips_globals(self):
        text = "interface GigabitEthernet0/1\n switchport mode access\n"
        rep = audit(text)
        assert "no-ntp" not in rep.codes()
        assert "access-vlan1" in rep.codes()
        assert rep.skipped and "фрагмент" in rep.skipped[0]

    def test_no_vty_noted(self):
        rep = audit("Building configuration...\nhostname X\nend\n")
        assert any("line vty" in s for s in rep.skipped)

    def test_po_members_not_double_reported(self):
        text = textwrap.dedent("""\
            interface GigabitEthernet0/1
             description up
             switchport mode trunk
             channel-group 1 mode active
            interface Port-channel1
             description up
             switchport mode trunk
        """)
        rep = audit(text)
        assert rep.get("trunk-all-vlans").where == ["Po1"]

    def test_routed_port_not_flagged_as_dtp(self):
        rep = audit("interface GigabitEthernet0/1\n description WAN\n no switchport\n ip address 1.1.1.1 255.255.255.0\n")
        assert "port-dtp" not in rep.codes()

    def test_router_l3_port_not_dtp_or_vlan1(self):
        text = "interface GigabitEthernet0/0/0\n no ip address\n negotiation auto\n!\n"
        rep = audit(text)
        assert "port-empty" not in rep.codes() and "port-dtp" not in rep.codes()
        assert rep.get("l3-port-idle").where == ["Gi0/0/0"]

    def test_unusable_vty_skipped(self):
        text = "Building configuration...\nhostname X\nline vty 5 15\n!\nend\n"
        rep = audit(text)
        assert not {"vty-transport-default", "vty-no-acl", "ssh-v1"} & rep.codes()
        assert any("line vty" in s for s in rep.skipped)

    def test_empty_vty_checked_with_aaa(self):
        text = "Building configuration...\nhostname X\naaa new-model\nline vty 0 4\n!\nend\n"
        assert "vty-transport-default" in audit(text).codes()

    def test_type7_password_not_plaintext(self):
        text = CISCO_CLEAN.replace("service password-encryption\n", "") + "line vty 0 15\n password 7 0822455D0A16\n"
        assert "password-encryption" not in audit(text).codes()

    def test_example_file(self):
        text = read_file_safe("examples/cisco-access-switch.txt")
        rep = audit(text)
        assert rep.get("access-no-bpduguard").where == ["Gi1/0/1-6, Gi1/0/9-10"]
        assert "trunk-all-vlans" not in rep.codes()   # Po1 має allowed
        assert "port-empty" not in rep.codes()        # Gi1/0/11-12 вимкнені


# ─────────────────────────────────────────────────────────────────────────────
#  FortiGate / FortiSwitch
# ─────────────────────────────────────────────────────────────────────────────

class TestFortiGateMessy:
    rep = audit(FGT_MESSY)

    def test_vendor(self):
        assert self.rep.vendor == "FortiGate"

    def test_expected_codes(self):
        expected = {"admin-cleartext", "admin-on-wan", "iface-no-description", "admin-no-trusthost",
                    "admin-default-name", "admin-no-2fa", "admin-timeout", "weak-crypto", "weak-tls",
                    "default-hostname", "no-ntp", "no-remote-log", "snmp-default-community",
                    "snmp-any-host", "no-password-policy", "policy-any-any", "policy-no-log"}
        assert expected <= self.rep.codes(), expected - self.rep.codes()

    def test_cleartext(self):
        assert self.rep.get("admin-cleartext").where == ["wan1: http", "lan: telnet"]
        assert "unselect allowaccess http telnet" in self.rep.get("admin-cleartext").fix

    def test_wan(self):
        assert self.rep.get("admin-on-wan").where == ["wan1: https, ssh, http"]

    def test_admins(self):
        assert self.rep.get("admin-no-trusthost").where == ["admin"]
        assert self.rep.get("admin-no-2fa").where == ["admin"]

    def test_disabled_policy_ignored(self):
        assert self.rep.get("policy-any-any").where == ["ID 1 «ANY-ANY»"]

    def test_no_description_only_unnamed(self):
        assert self.rep.get("iface-no-description").where == ["wan1", "lan"]


class TestFortiGateClean:
    def test_no_findings(self):
        rep = audit(FGT_CLEAN)
        assert rep.findings == [], [f.code for f in rep.findings]

    def test_example_is_fragment(self):
        rep = audit(read_file_safe("examples/fortigate.conf"))
        assert rep.skipped and "system admin" in rep.skipped[0]
        assert "admin-timeout" in rep.codes()
        assert "admin-on-wan" not in rep.codes()   # на wan1 лише ping


class TestFortiSwitch:
    rep = audit(FSW)

    def test_vendor(self):
        assert self.rep.vendor == "FortiSwitch"

    def test_ports(self):
        assert self.rep.get("sw-allowed-all").where == ["port2"]
        assert self.rep.get("sw-stp-off").where == ["port2"]
        assert self.rep.get("sw-native1").where == ["port2"]
        assert self.rep.get("no-description").where == ["port2"]


# ─────────────────────────────────────────────────────────────────────────────
#  DOCX-звіт
# ─────────────────────────────────────────────────────────────────────────────

class TestAuditDocx:
    def test_report_table(self, tmp_path):
        from docx import Document
        from main import build_audit_docx
        parsed = parse_config(CISCO_MESSY)
        rep = run_audit(CISCO_MESSY, parsed)
        out = tmp_path / "a.docx"
        build_audit_docx(rep, parsed, {"model": "WS-C2960X", "serial": "FOC1"}, str(out))
        doc = Document(str(out))
        assert doc.tables[0].cell(0, 0).text == "Switch — аудит конфігурації"
        table = doc.tables[-1]
        assert len(table.rows) == len(rep.findings) + 1
        assert table.cell(1, 0).text == "КРИТИЧНО"
        first = rep.findings[0]
        assert table.cell(1, 1).paragraphs[0].text == first.title
        assert table.cell(1, 1).paragraphs[1].text == first.why
        fix_runs = [r for p in table.cell(1, 3).paragraphs for r in p.runs]
        assert fix_runs and all(r._r.rPr.rFonts.get(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii") == "Consolas" for r in fix_runs)
        body = "\n".join(p.text for p in doc.paragraphs)
        assert "Cisco IOS:" in body and "критичн" in body

    def test_clean_report(self, tmp_path):
        from docx import Document
        from main import build_audit_docx
        parsed = parse_config(CISCO_CLEAN)
        out = tmp_path / "b.docx"
        build_audit_docx(run_audit(CISCO_CLEAN, parsed), parsed, {}, str(out))
        doc = Document(str(out))
        assert "Недолаштованого не знайдено." in [p.text for p in doc.paragraphs]
        assert len(doc.tables) == 2   # шапка + інфо, без таблиці знахідок
