#!/usr/bin/env python3
"""
network-documentr — Cisco 'show running-config' / Fortinet 'show full-configuration'
→ DOCX generator
"""

import re
import os
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from audit import run_audit, report_text, summary_text, short_summary, SEV_LABEL

# ── Colours ───────────────────────────────────────────────────────────────────
C_DARK_BLUE   = "1A2B4A"
C_MID_BLUE    = "2C4A7C"
C_ACCENT_BLUE = "3B6BA5"
C_LIGHT_BLUE  = "E8EFF8"
C_ROW_ALT     = "F5F8FC"
C_PO_ROW      = "EEF2FF"
C_VLAN_ROW    = "F0FDF4"
C_VLAN_DEAD   = "ECECEC"
C_BORDER      = "CBD5E1"
C_WHITE       = "FFFFFF"

RGB_WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
RGB_DARK_TEXT  = RGBColor(0x1E, 0x29, 0x3B)
RGB_GREY_LABEL = RGBColor(0x94, 0xA3, 0xB8)
RGB_AMBER      = RGBColor(0xB4, 0x53, 0x09)
RGB_DEAD_TEXT  = RGBColor(0xB0, 0xB0, 0xB0)

# A4 portrait with 14 mm margins → 182 mm usable width
USABLE_MM = 182.0

# ── Interface classification ──────────────────────────────────────────────────
PHYSICAL_PREFIXES = (
    "GigabitEthernet", "FastEthernet", "TenGigabitEthernet",
    "TwentyFiveGigE", "FortyGigabitEthernet", "HundredGigE", "Management",
)
LOGICAL_PREFIXES = ("Port-channel", "port-channel", "Vlan", "Loopback")

# ═════════════════════════════════════════════════════════════════════════════
#  PARSER
# ═════════════════════════════════════════════════════════════════════════════

def read_file_safe(path: str) -> str:
    for enc in ("utf-8", "cp1251", "latin-1"):
        try:
            with open(path, encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    with open(path, encoding="utf-8", errors="replace") as f:
        return f.read()


def short_iface(name: str) -> str:
    """GigabitEthernet1/0/1 → Gi1/0/1, etc."""
    for long, short in [
        ("GigabitEthernet",     "Gi"),
        ("FastEthernet",        "Fa"),
        ("TenGigabitEthernet",  "Te"),
        ("TwentyFiveGigE",      "Twe"),
        ("FortyGigabitEthernet","Fo"),
        ("HundredGigE",         "Hu"),
        ("Port-channel",        "Po"),
        ("port-channel",        "Po"),
        ("Vlan",                "Vlan"),
        ("Loopback",            "Lo"),
        ("Management",          "Mgmt"),
    ]:
        if name.startswith(long):
            return short + name[len(long):]
    return name


def _blank_iface(name: str) -> dict:
    """Interface record with every field the DOCX builder expects."""
    return {
        "name":            name,
        "description":     "",
        "mode":            "",
        "access_vlan":     "",
        "native_vlan":     "",
        "trunk_vlans":     "",
        "ip_address":      "",
        "shutdown":        False,
        "notes":           "",
        "iface_class":     "physical",
        "row_kind":        "",            # "" | "po" | "vlan" | "lo" — colour/counter hint
        "voice_vlan":              "",
        "stp_portfast":            False,
        "stp_rootguard":           False,
        "stp_bpduguard":           False,
        "port_security":           False,
        "port_security_max":       "",
        "port_security_violation": "",
        "port_security_macs":      [],
        "port_security_sticky":    False,
    }


def _collect_used_vlans(interfaces: list) -> tuple:
    """Return (used_ids: set[str], all_used: bool).
    used_ids — VLAN IDs referenced on any interface.
    all_used  — True if any trunk allows 'all' or an unexpandable wide range.
    """
    used = set()
    all_used = False
    for iface in interfaces:
        for field in ("access_vlan", "native_vlan", "voice_vlan"):
            val = iface.get(field, "")
            if val:
                used.add(val)
        trunk = iface.get("trunk_vlans", "")
        if trunk:
            for part in trunk.split(","):
                part = part.strip()
                if not part or part.lower() == "none":
                    continue
                if part.lower() == "all":
                    all_used = True
                    continue
                m = re.match(r"^(\d+)-(\d+)$", part)
                if m:
                    lo, hi = int(m.group(1)), int(m.group(2))
                    if hi - lo <= 1000:
                        used.update(str(i) for i in range(lo, hi + 1))
                    else:
                        all_used = True
                elif re.match(r"^\d+$", part):
                    used.add(part)
    return used, all_used


def parse_showrun(text: str) -> dict:
    result = {"hostname": "", "mgmt_ip": "", "interfaces": []}

    m = re.search(r"^hostname\s+(\S+)", text, re.MULTILINE)
    if m:
        result["hostname"] = m.group(1)

    m = re.search(r"^ip ssh source-interface\s+(\S+)", text, re.MULTILINE)
    result["ssh_source_iface"] = m.group(1) if m else ""

    iface_blocks = re.split(r"(?=^interface\s)", text, flags=re.MULTILINE)

    for block in iface_blocks:
        if not block.strip().startswith("interface"):
            continue

        lines = block.strip().splitlines()
        iface_name = lines[0].split(None, 1)[1].strip()

        iface = _blank_iface(iface_name)

        for line in lines[1:]:
            line = line.strip()

            if line.startswith("description "):
                iface["description"] = line[len("description "):].strip()

            elif line == "switchport mode access":
                iface["mode"] = "access"

            elif line == "switchport mode trunk":
                iface["mode"] = "trunk"

            elif re.match(r"no switchport", line):
                iface["mode"] = "routed"

            elif re.match(r"switchport access vlan (\d+)", line):
                iface["access_vlan"] = re.search(r"\d+", line).group()

            elif re.match(r"switchport trunk native vlan (\d+)", line):
                iface["native_vlan"] = re.search(r"\d+", line).group()

            elif re.match(r"switchport trunk allowed vlan", line):
                vlans = re.sub(r"switchport trunk allowed vlan\s*(add\s*)?", "", line).strip()
                if iface["trunk_vlans"]:
                    iface["trunk_vlans"] += "," + vlans
                else:
                    iface["trunk_vlans"] = vlans

            elif re.match(r"ip address ", line):
                iface["ip_address"] = line[len("ip address "):].strip()

            elif line == "shutdown":
                iface["shutdown"] = True

            elif re.match(r"channel-group (\d+)", line):
                po_num = re.search(r"channel-group (\d+)", line).group(1)
                iface["notes"] = f"Po{po_num} member"

            elif re.match(r"switchport voice vlan (\d+)", line):
                iface["voice_vlan"] = re.search(r"switchport voice vlan (\d+)", line).group(1)

            elif re.match(r"spanning-tree portfast", line):
                iface["stp_portfast"] = True

            elif line == "spanning-tree guard root":
                iface["stp_rootguard"] = True

            elif line == "spanning-tree bpduguard enable":
                iface["stp_bpduguard"] = True

            elif re.match(r"switchport port-security maximum (\d+)", line):
                iface["port_security_max"] = re.search(r"switchport port-security maximum (\d+)", line).group(1)
                iface["port_security"] = True

            elif re.match(r"switchport port-security violation (\S+)", line):
                iface["port_security_violation"] = re.search(r"switchport port-security violation (\S+)", line).group(1)
                iface["port_security"] = True

            elif re.match(r"switchport port-security mac-address sticky", line):
                iface["port_security_sticky"] = True
                iface["port_security"] = True
                m_sticky = re.search(r"switchport port-security mac-address sticky\s+(\S+)", line)
                if m_sticky:
                    iface["port_security_macs"].append(m_sticky.group(1))

            elif re.match(r"switchport port-security mac-address (\S+)", line):
                mac = re.search(r"switchport port-security mac-address (\S+)", line).group(1)
                iface["port_security_macs"].append(mac)
                iface["port_security"] = True

            elif line == "switchport port-security":
                iface["port_security"] = True

        # Build notes: Po member → Voice:X → PortFast → RootGuard → BPDUGuard → PortSec … → shutdown
        notes_parts = [iface["notes"]] if iface["notes"] else []
        if iface["voice_vlan"]:   notes_parts.append(f"Voice:{iface['voice_vlan']}")
        if iface["stp_portfast"]: notes_parts.append("PortFast")
        if iface["stp_rootguard"]:notes_parts.append("RootGuard")
        if iface["stp_bpduguard"]:notes_parts.append("BPDUGuard")
        if iface["port_security"]:
            ps_parts = ["PortSec"]
            if iface["port_security_max"]:       ps_parts.append(f"max:{iface['port_security_max']}")
            if iface["port_security_violation"]: ps_parts.append(f"violation:{iface['port_security_violation']}")
            if iface["port_security_sticky"]:    ps_parts.append("sticky")
            notes_parts.append(" ".join(ps_parts))
            for mac in iface["port_security_macs"]:
                notes_parts.append(f"MAC:{mac}")
        if iface["shutdown"]:     notes_parts.append("shutdown")
        iface["notes"] = "\n".join(notes_parts)

        # Mode fallback
        if not iface["mode"]:
            if iface["ip_address"] or iface_name.lower().startswith("vlan"):
                iface["mode"] = "routed"
            elif any(iface_name.lower().startswith(p.lower())
                     for p in ("port-channel",)):
                iface["mode"] = "trunk"

        # Classify physical vs logical
        if any(iface_name.startswith(p) for p in PHYSICAL_PREFIXES):
            iface["iface_class"] = "physical"
        elif any(iface_name.startswith(p) for p in LOGICAL_PREFIXES):
            iface["iface_class"] = "logical"
        else:
            iface["iface_class"] = "physical"

        if iface_name.lower().startswith("port-channel"):
            iface["row_kind"] = "po"
        elif iface_name.startswith("Vlan"):
            iface["row_kind"] = "vlan"
        elif iface_name.startswith("Loopback"):
            iface["row_kind"] = "lo"

        result["interfaces"].append(iface)

    # Parse VLAN definitions (single-VID blocks only; skip ranges, "vlan database")
    result["vlans"] = []
    vlan_blocks = re.split(r"(?=^vlan\s)", text, flags=re.MULTILINE)
    for vblock in vlan_blocks:
        lines_v = vblock.strip().splitlines()
        if not lines_v:
            continue
        m_vid = re.match(r"^vlan\s+(\d+)$", lines_v[0].strip())
        if not m_vid:
            continue
        vlan_name = ""
        for vl in lines_v[1:]:
            if vl.strip().startswith("name "):
                vlan_name = vl.strip()[len("name "):].strip()
                break
        result["vlans"].append({"id": m_vid.group(1), "name": vlan_name})

    # Mark each VLAN as used (referenced on any port) or dead
    used_ids, trunk_all = _collect_used_vlans(result["interfaces"])
    for vlan in result["vlans"]:
        vlan["used"] = trunk_all or vlan["id"] in used_ids
    result["vlans"].sort(key=lambda v: int(v["id"]))

    # Resolve MGMT IP: prefer SSH source-interface, fallback to first IP found
    ssh_src = result.get("ssh_source_iface", "")
    for iface in result["interfaces"]:
        if ssh_src and iface["name"] == ssh_src and iface["ip_address"]:
            result["mgmt_ip"] = iface["ip_address"].split()[0]
            break
    if not result["mgmt_ip"]:
        for iface in result["interfaces"]:
            if iface["ip_address"]:
                result["mgmt_ip"] = iface["ip_address"].split()[0]
                break

    return result


# ═════════════════════════════════════════════════════════════════════════════
#  FORTINET PARSER (FortiGate / FortiSwitch — 'show full-configuration')
# ═════════════════════════════════════════════════════════════════════════════

def detect_vendor(text: str) -> str:
    """Return 'fortinet' for FortiOS/FortiSwitchOS configs, otherwise 'cisco'."""
    if re.search(r"^#config-version=", text, re.MULTILINE):
        return "fortinet"
    if re.search(r"^\s*config system (global|interface)\s*$", text, re.MULTILINE):
        return "fortinet"
    if re.search(r"^\s*config switch (interface|physical-port)\s*$", text, re.MULTILINE):
        return "fortinet"
    return "cisco"


def _forti_unquote(val: str) -> str:
    val = val.strip()
    if len(val) >= 2 and val[0] == '"' and val[-1] == '"':
        val = val[1:-1]
    return val


def _forti_values(val: str) -> list:
    """'\"port51\" \"port52\"' → ['port51', 'port52']; unquoted → split on spaces."""
    vals = re.findall(r'"([^"]*)"', val)
    return vals if vals else val.split()


def _forti_section(text: str, path: str) -> list:
    """Entries of FortiOS 'config <path>' … 'end' sections as (edit_name, {set key: raw value}).
    Nested 'config … end' blocks inside an edit are skipped; multiple sections
    with the same path (multi-VDOM) are merged. Returns [] if absent."""
    entries = []
    in_section = False
    depth = 0
    name = None
    settings = None
    target = "config " + path
    for raw in text.splitlines():
        s = raw.strip()
        if not in_section:
            if s == target:
                in_section = True
                depth = 0
                name, settings = None, None
            continue
        if s.startswith("config "):
            depth += 1
        elif s == "end":
            if depth == 0:
                in_section = False
            else:
                depth -= 1
        elif depth == 0:
            if s.startswith("edit "):
                name = _forti_unquote(s[len("edit "):])
                settings = {}
            elif s == "next":
                if name is not None:
                    entries.append((name, settings))
                name, settings = None, None
            elif s.startswith("set ") and settings is not None:
                kv = s[len("set "):].split(None, 1)
                if kv:
                    settings[kv[0]] = kv[1].strip() if len(kv) > 1 else ""
    return entries


def _forti_hostname(text: str) -> str:
    m = re.search(r"^\s*set hostname\s+(.+?)\s*$", text, re.MULTILINE)
    return _forti_unquote(m.group(1)) if m else ""


def _ip_with_prefix(ip_field: str) -> str:
    """'203.0.113.2 255.255.255.252' → '203.0.113.2/30'."""
    parts = ip_field.split()
    if len(parts) == 2:
        try:
            prefix = sum(bin(int(o)).count("1") for o in parts[1].split("."))
            return f"{parts[0]}/{prefix}"
        except ValueError:
            pass
    return parts[0] if parts else ""


def _forti_pick_mgmt(candidates: list) -> str:
    """candidates: (name, allowaccess, ip). Priority: 'mgmt' iface →
    allowaccess з ssh/https → перший IP."""
    for n, _a, ip in candidates:
        if n.lower() in ("mgmt", "mgmt1", "management"):
            return ip
    for _n, a, ip in candidates:
        if "ssh" in a or "https" in a:
            return ip
    return candidates[0][2] if candidates else ""


def _forti_fill_switchport(iface: dict, ss: dict):
    """Map FortiSwitch 'config switch interface' settings onto access/trunk fields."""
    native  = ss.get("native-vlan", "")
    allowed = ss.get("allowed-vlans", "").strip()
    allowed = re.sub(r"[\s,]+", ",", allowed) if allowed else ""
    if ss.get("allowed-vlans-all", "") in ("enable", "enabled"):
        allowed = "all"
    if allowed:
        iface["mode"] = "trunk"
        iface["trunk_vlans"] = allowed
        iface["native_vlan"] = native
    elif native:
        iface["mode"] = "access"
        iface["access_vlan"] = native


def _forti_parse_gate(result: dict, sys_iface: list):
    """FortiGate: усі інтерфейси в 'config system interface'."""
    member_of = {}
    for name, s in sys_iface:
        if s.get("type", "") in ("aggregate", "redundant", "switch",
                                 "hard-switch", "software-switch"):
            for mem in _forti_values(s.get("member", "")):
                member_of[mem] = name

    mgmt_candidates = []
    for name, s in sys_iface:
        iface  = _blank_iface(name)
        typ    = s.get("type", "")
        vlanid = s.get("vlanid", "")
        ip     = s.get("ip", "")
        if ip and not ip.startswith("0.0.0.0"):
            iface["ip_address"] = ip
        alias = _forti_unquote(s.get("alias", ""))
        descr = _forti_unquote(s.get("description", ""))
        iface["description"] = " / ".join(p for p in (alias, descr) if p)
        iface["shutdown"] = s.get("status", "").strip() == "down"

        notes = []
        if vlanid and not typ:
            typ = "vlan"
        if typ == "vlan":
            iface["iface_class"] = "logical"
            iface["row_kind"]    = "vlan"
            iface["access_vlan"] = vlanid
            parent = _forti_unquote(s.get("interface", ""))
            notes.append(f"VLAN {vlanid}" + (f" @ {parent}" if parent else ""))
        elif typ in ("aggregate", "redundant"):
            iface["iface_class"] = "logical"
            iface["row_kind"]    = "po"
            iface["mode"]        = "lacp" if typ == "aggregate" else "redundant"
            members = _forti_values(s.get("member", ""))
            if members:
                notes.append("members: " + ", ".join(members))
        elif typ == "loopback":
            iface["iface_class"] = "logical"
            iface["row_kind"]    = "lo"
        elif typ == "tunnel":
            iface["iface_class"] = "logical"
            notes.append("tunnel")
        elif typ in ("switch", "hard-switch", "software-switch"):
            iface["iface_class"] = "logical"
            iface["mode"]        = "switch"
            members = _forti_values(s.get("member", ""))
            if members:
                notes.append("members: " + ", ".join(members))

        if not iface["mode"]:
            if s.get("mode", "") == "dhcp":
                iface["mode"] = "dhcp"
            elif iface["ip_address"]:
                iface["mode"] = "routed"

        # Physical table has no IP column — surface the address in notes
        if iface["iface_class"] == "physical" and iface["ip_address"]:
            notes.insert(0, f"IP: {_ip_with_prefix(iface['ip_address'])}")

        if name in member_of:
            notes.append(f"{member_of[name]} member")
        role = s.get("role", "")
        if role and role != "undefined":
            notes.append(f"role: {role}")
        if s.get("dedicated-to", "") == "management":
            notes.append("dedicated mgmt")
        allow = s.get("allowaccess", "")
        if allow:
            notes.append(allow)
        if iface["shutdown"]:
            notes.append("shutdown")
        iface["notes"] = "\n".join(notes)

        result["interfaces"].append(iface)
        if iface["ip_address"]:
            mgmt_candidates.append((name, allow, iface["ip_address"].split()[0]))

    result["mgmt_ip"] = _forti_pick_mgmt(mgmt_candidates)


def _forti_parse_switch(result: dict, sw_phys: list, sw_iface: list,
                        sw_trunk: list, sw_vlan: list, sys_iface: list):
    """Standalone FortiSwitch: порти з 'config switch …', mgmt — з 'config system interface'."""
    phys = dict(sw_phys)
    swif = dict(sw_iface)

    member_of   = {}
    trunk_names = set()
    for tname, ts in sw_trunk:
        trunk_names.add(tname)
        for mem in _forti_values(ts.get("members", "")):
            member_of[mem] = tname

    sys_names = {n for n, _s in sys_iface}

    # Physical ports: order from physical-port section, then any extra
    # switch-interface entries that are not trunks / system interfaces
    port_order = [n for n, _s in sw_phys]
    for n, _s in sw_iface:
        if n not in phys and n not in trunk_names and n not in sys_names:
            port_order.append(n)

    for name in port_order:
        ps = phys.get(name, {})
        ss = swif.get(name, {})
        iface = _blank_iface(name)
        iface["description"] = _forti_unquote(
            ps.get("description", "") or ss.get("description", ""))
        iface["shutdown"] = ps.get("status", "").strip() == "down"
        _forti_fill_switchport(iface, ss)

        notes = []
        if name in member_of:
            notes.append(f"{member_of[name]} member")
            if not iface["mode"]:
                iface["mode"] = "lag"
        sec = ss.get("security-mode", "")
        if sec and sec != "none":
            iface["port_security"] = True
            notes.append(sec)
        if ss.get("stp-state", "") == "disabled":
            notes.append("STP off")
        if iface["shutdown"]:
            notes.append("shutdown")
        iface["notes"] = "\n".join(notes)
        result["interfaces"].append(iface)

    # Link aggregation trunks → logical (Po analogue)
    for tname, ts in sw_trunk:
        iface = _blank_iface(tname)
        iface["iface_class"] = "logical"
        iface["row_kind"]    = "po"
        iface["description"] = _forti_unquote(ts.get("description", ""))
        _forti_fill_switchport(iface, swif.get(tname, {}))

        notes = []
        members = _forti_values(ts.get("members", ""))
        if members:
            notes.append("members: " + ", ".join(members))
        if iface["trunk_vlans"]:
            notes.append(f"VLANs: {iface['trunk_vlans']}")
            if iface["native_vlan"]:
                notes.append(f"native: {iface['native_vlan']}")
        elif iface["access_vlan"]:
            notes.append(f"VLAN: {iface['access_vlan']}")
        iface["mode"] = ts.get("mode", "") or "static"
        iface["notes"] = "\n".join(notes)
        result["interfaces"].append(iface)

    # System interfaces (mgmt / internal) → logical, carry the IP
    mgmt_candidates = []
    for name, s in sys_iface:
        iface = _blank_iface(name)
        iface["iface_class"] = "logical"
        if s.get("type", "") == "vlan" or s.get("vlanid", ""):
            iface["row_kind"] = "vlan"
        ip = s.get("ip", "")
        if ip and not ip.startswith("0.0.0.0"):
            iface["ip_address"] = ip
        iface["description"] = _forti_unquote(s.get("description", ""))
        iface["shutdown"] = s.get("status", "").strip() == "down"
        if iface["ip_address"]:
            iface["mode"] = "routed"

        notes = []
        vlanid = s.get("vlanid", "")
        if vlanid:
            notes.append(f"VLAN {vlanid}")
        allow = s.get("allowaccess", "")
        if allow:
            notes.append(allow)
        if iface["shutdown"]:
            notes.append("shutdown")
        iface["notes"] = "\n".join(notes)
        result["interfaces"].append(iface)
        if iface["ip_address"]:
            mgmt_candidates.append((name, allow, iface["ip_address"].split()[0]))

    result["mgmt_ip"] = _forti_pick_mgmt(mgmt_candidates)

    # VLAN definitions from 'config switch vlan'
    for vid, vs in sw_vlan:
        if re.match(r"^\d+$", vid):
            result["vlans"].append(
                {"id": vid, "name": _forti_unquote(vs.get("description", ""))})


def parse_fortinet(text: str) -> dict:
    """Parse FortiGate / FortiSwitch 'show full-configuration' into the same
    schema as parse_showrun()."""
    result = {"hostname": _forti_hostname(text), "mgmt_ip": "",
              "ssh_source_iface": "", "interfaces": [], "vlans": []}

    sys_iface = _forti_section(text, "system interface")
    sw_phys   = _forti_section(text, "switch physical-port")
    sw_iface  = _forti_section(text, "switch interface")
    sw_trunk  = _forti_section(text, "switch trunk")
    sw_vlan   = _forti_section(text, "switch vlan")

    if sw_iface or sw_phys:
        _forti_parse_switch(result, sw_phys, sw_iface, sw_trunk, sw_vlan, sys_iface)
    else:
        _forti_parse_gate(result, sys_iface)

    used_ids, trunk_all = _collect_used_vlans(result["interfaces"])
    for vlan in result["vlans"]:
        vlan["used"] = trunk_all or vlan["id"] in used_ids
    result["vlans"].sort(key=lambda v: int(v["id"]))

    return result


def parse_config(text: str) -> dict:
    """Vendor-aware dispatcher: Cisco IOS або FortiOS/FortiSwitchOS."""
    if detect_vendor(text) == "fortinet":
        return parse_fortinet(text)
    return parse_showrun(text)


def parse_version_inventory(text: str) -> dict:
    """Extract model and serial from show inventory / show version blocks."""
    model = ""
    serial = ""

    # Priority 1: show inventory — first PID line (chassis / main unit)
    m = re.search(r"^PID:\s*(\S+)\s*,.*SN:\s*(\S+)", text, re.MULTILINE)
    if m:
        model = m.group(1)
        serial = m.group(2)

    # Priority 2: show version — Model number
    if not model:
        m = re.search(r"^Model [Nn]umber\s*:\s*(\S+)", text, re.MULTILINE)
        if m:
            model = m.group(1)

    # Priority 3: show version — System serial number
    if not serial:
        m = re.search(r"^System [Ss]erial [Nn]umber\s*:\s*(\S+)", text, re.MULTILINE)
        if m:
            serial = m.group(1)

    # Priority 4: show version — Processor board ID (older platforms e.g. 3750)
    if not serial:
        m = re.search(r"^Processor board ID\s+(\S+)", text, re.MULTILINE)
        if m:
            serial = m.group(1)

    # Fortinet: model from the '#config-version=FGT60E-6.4.5-…' header
    if not model:
        m = re.search(r"^#config-version=(\w+)-", text, re.MULTILINE)
        if m:
            model = m.group(1)

    # Fortinet: serial from a pasted 'get system status' block
    if not serial:
        m = re.search(r"^Serial-Number:\s*(\S+)", text, re.MULTILINE)
        if m:
            serial = m.group(1)

    return {"model": model, "serial": serial}


def classify_interfaces(interfaces: list):
    physical = [i for i in interfaces if i["iface_class"] == "physical"]
    logical  = [i for i in interfaces if i["iface_class"] == "logical"]
    return physical, logical


def format_load_status(parsed: dict) -> str:
    """Build status bar string shown after a config file is loaded."""
    physical, logical = classify_interfaces(parsed["interfaces"])

    def _kind(i):
        rk = i.get("row_kind", "")
        if rk:
            return rk
        n = i.get("name", "")
        if n.lower().startswith("port-channel"):
            return "po"
        if n.startswith("Vlan"):
            return "vlan"
        if n.startswith("Loopback"):
            return "lo"
        return ""

    po    = sum(1 for i in logical if _kind(i) == "po")
    vlans = sum(1 for i in logical if _kind(i) == "vlan")
    loops = sum(1 for i in logical if _kind(i) == "lo")
    parts = []
    if parsed.get("mgmt_ip"):
        parts.append(f"MGMT: {parsed['mgmt_ip']}")
    parts.append(f"Порти: {len(physical)}")
    if po:    parts.append(f"Po: {po}")
    if vlans: parts.append(f"VLAN: {vlans}")
    if loops: parts.append(f"Lo: {loops}")
    return "  ·  ".join(parts)


# ═════════════════════════════════════════════════════════════════════════════
#  XML HELPERS
# ═════════════════════════════════════════════════════════════════════════════

def set_cell_bg(cell, hex_color: str):
    """Apply solid background fill to a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    for el in tcPr.findall(qn("w:shd")):
        tcPr.remove(el)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)


def set_table_borders(table, color_hex: str = "CBD5E1",
                      sz: str = "4", val: str = "single"):
    """Set all table borders."""
    tblPr = table._tbl.tblPr
    for el in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(el)
    tblBorders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"),   val)
        b.set(qn("w:sz"),    sz)
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), color_hex)
        tblBorders.append(b)
    tblPr.append(tblBorders)


def set_table_fixed_layout(table):
    """Force fixed column widths (no autofit)."""
    tblPr = table._tbl.tblPr
    for el in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(el)
    el = OxmlElement("w:tblLayout")
    el.set(qn("w:type"), "fixed")
    tblPr.append(el)


def set_table_width(table, width_mm: float):
    """Set total table width."""
    tblPr = table._tbl.tblPr
    for el in tblPr.findall(qn("w:tblW")):
        tblPr.remove(el)
    el = OxmlElement("w:tblW")
    el.set(qn("w:w"),    str(int(width_mm * 56.7)))
    el.set(qn("w:type"), "dxa")
    tblPr.append(el)


def set_cell_margins_table(table, top=40, left=60, bottom=40, right=60):
    """Set default cell padding for the whole table (twips)."""
    tblPr = table._tbl.tblPr
    for el in tblPr.findall(qn("w:tblCellMar")):
        tblPr.remove(el)
    mar = OxmlElement("w:tblCellMar")
    for side, val in [("top", top), ("left", left),
                      ("bottom", bottom), ("right", right)]:
        n = OxmlElement(f"w:{side}")
        n.set(qn("w:w"),    str(val))
        n.set(qn("w:type"), "dxa")
        mar.append(n)
    tblPr.append(mar)


def set_cell_margins(cell, top=60, left=80, bottom=60, right=60):
    """Set padding on a single cell (twips)."""
    tcPr  = cell._tc.get_or_add_tcPr()
    for el in tcPr.findall(qn("w:tcMar")):
        tcPr.remove(el)
    mar = OxmlElement("w:tcMar")
    for side, val in [("top", top), ("left", left),
                      ("bottom", bottom), ("right", right)]:
        n = OxmlElement(f"w:{side}")
        n.set(qn("w:w"),    str(val))
        n.set(qn("w:type"), "dxa")
        mar.append(n)
    tcPr.append(mar)


def set_repeat_header_row(row):
    """Mark row to repeat as header on each page."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))


def set_row_min_height(row, height_twips: int = 240):
    """Set minimum row height (240 twips ≈ 4.2 mm)."""
    trPr = row._tr.get_or_add_trPr()
    for el in trPr.findall(qn("w:trHeight")):
        trPr.remove(el)
    el = OxmlElement("w:trHeight")
    el.set(qn("w:val"),   str(height_twips))
    el.set(qn("w:hRule"), "atLeast")
    trPr.append(el)


def _set_run_font(run, font_name: str = "Calibri"):
    """Set font with full Cyrillic support (w:cs attribute)."""
    rPr    = run._r.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"),    font_name)
    rFonts.set(qn("w:hAnsi"),    font_name)
    rFonts.set(qn("w:cs"),       font_name)
    rFonts.set(qn("w:eastAsia"), font_name)
    existing = rPr.find(qn("w:rFonts"))
    if existing is not None:
        rPr.remove(existing)
    rPr.insert(0, rFonts)


def _write_cell(cell, text: str, bold: bool = False,
                color: RGBColor = None, font_size: float = 7,
                align=WD_ALIGN_PARAGRAPH.LEFT):
    """Write text into the first paragraph of a cell with full styling.
    '\\n' in text produces in-cell line breaks (w:br)."""
    para = cell.paragraphs[0]
    para.clear()
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after  = Pt(0)
    para.alignment = align
    parts = (text or "").split("\n")
    for i, part in enumerate(parts):
        run = para.add_run(part)
        run.bold = bold
        run.font.size = Pt(font_size)
        if color:
            run.font.color.rgb = color
        _set_run_font(run)
        if i < len(parts) - 1:
            br = OxmlElement("w:br")
            run._r.append(br)


# ═════════════════════════════════════════════════════════════════════════════
#  DOCX BUILDERS
# ═════════════════════════════════════════════════════════════════════════════

def _add_header_block(doc, hostname: str):
    """Dark-blue banner with hostname."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_fixed_layout(tbl)
    set_table_width(tbl, USABLE_MM)
    set_table_borders(tbl, val="none")

    cell = tbl.cell(0, 0)
    cell.width = Mm(USABLE_MM)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_bg(cell, C_DARK_BLUE)
    set_cell_margins(cell, top=120, left=120, bottom=120, right=120)
    _write_cell(cell, hostname, bold=True, color=RGB_WHITE, font_size=16)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after  = Pt(2)


def _add_info_block(doc, model: str, serial: str, mgmt_ip: str, updated: str):
    """Light-blue info row: Model | Serial | MGMT IP | Updated."""
    COL_W  = USABLE_MM / 4
    labels = ["МОДЕЛЬ / ВЕНДОР", "СЕРІЙНИЙ НОМЕР", "MGMT IP-АДРЕСА", "ОНОВЛЕНО"]
    values = [model or "—", serial or "—", mgmt_ip or "—", updated]

    tbl = doc.add_table(rows=1, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_fixed_layout(tbl)
    set_table_width(tbl, USABLE_MM)
    set_table_borders(tbl, color_hex=C_BORDER, sz="4")

    for ci, (lbl, val) in enumerate(zip(labels, values)):
        cell = tbl.cell(0, ci)
        cell.width = Mm(COL_W)
        set_cell_bg(cell, C_LIGHT_BLUE)
        set_cell_margins(cell, top=60, left=80, bottom=60, right=40)

        # label — first paragraph already present in cell
        lp = cell.paragraphs[0]
        lp.clear()
        lp.paragraph_format.space_before = Pt(0)
        lp.paragraph_format.space_after  = Pt(1)
        lr = lp.add_run(lbl)
        lr.font.size = Pt(8)
        lr.font.color.rgb = RGB_GREY_LABEL
        _set_run_font(lr)

        # value — second paragraph
        vp = cell.add_paragraph()
        vp.paragraph_format.space_before = Pt(0)
        vp.paragraph_format.space_after  = Pt(0)
        vr = vp.add_run(val)
        vr.bold = True
        vr.font.size = Pt(10)
        vr.font.color.rgb = RGB_DARK_TEXT
        _set_run_font(vr)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after  = Pt(2)


def _add_section_header(doc, text: str):
    """Mid-blue section label bar."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_fixed_layout(tbl)
    set_table_width(tbl, USABLE_MM)
    set_table_borders(tbl, val="none")

    cell = tbl.cell(0, 0)
    cell.width = Mm(USABLE_MM)
    set_cell_bg(cell, C_MID_BLUE)
    set_cell_margins(cell, top=60, left=100, bottom=60, right=60)
    _write_cell(cell, text, bold=True, color=RGB_WHITE, font_size=10)


def _add_interface_table(doc, interfaces: list, logical: bool = False):
    """
    Compact interface table.
    Physical (logical=False): ПОРТ | ОПИС | РЕЖИМ | ACC/NAT VLAN | TRUNK VLANs | ПРИМІТКА
    Logical  (logical=True):  ПОРТ | ОПИС | РЕЖИМ | IP-АДРЕСА | ПРИМІТКА
    """
    if not interfaces:
        sp = doc.add_paragraph("(немає інтерфейсів)")
        sp.paragraph_format.space_before = Pt(2)
        sp.paragraph_format.space_after  = Pt(4)
        return

    if logical:
        # 5 cols, sum = 182 mm
        COL_W_MM = [18, 82, 22, 37, 23]
        COL_HDR  = ["ПОРТ", "ОПИС / ПРИЗНАЧЕННЯ", "РЕЖИМ", "IP-АДРЕСА", "ПРИМІТКА"]
        CENTER   = {0, 2, 3, 4}
    else:
        # 6 cols, sum = 182 mm
        COL_W_MM = [18, 45, 22, 32, 42, 23]
        COL_HDR  = ["ПОРТ", "ОПИС / ПРИЗНАЧЕННЯ", "РЕЖИМ",
                    "ACC/NAT VLAN", "TRUNK VLANs", "ПРИМІТКА"]
        CENTER   = {0, 2, 3, 4, 5}

    n_cols = len(COL_W_MM)
    tbl = doc.add_table(rows=1, cols=n_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_fixed_layout(tbl)
    set_table_width(tbl, USABLE_MM)
    set_table_borders(tbl, color_hex=C_BORDER, sz="4")
    set_cell_margins_table(tbl, top=15, left=60, bottom=15, right=40)

    # ── Header row ────────────────────────────────────────────────────────
    hdr = tbl.rows[0]
    set_repeat_header_row(hdr)
    set_row_min_height(hdr, 400)

    for ci, (h, w) in enumerate(zip(COL_HDR, COL_W_MM)):
        c = hdr.cells[ci]
        c.width = Mm(w)
        set_cell_bg(c, C_ACCENT_BLUE)
        _write_cell(c, h, bold=True, color=RGB_WHITE, font_size=10,
                    align=WD_ALIGN_PARAGRAPH.CENTER if ci in CENTER
                    else WD_ALIGN_PARAGRAPH.LEFT)

    # ── Data rows ─────────────────────────────────────────────────────────
    for ri, iface in enumerate(interfaces):
        short = short_iface(iface["name"])

        kind = iface.get("row_kind", "")
        if kind == "po" or short.startswith("Po"):
            bg = C_PO_ROW
        elif kind in ("vlan", "lo") or short.startswith("Vlan") or short.startswith("Lo"):
            bg = C_VLAN_ROW
        elif ri % 2 == 0:
            bg = C_WHITE
        else:
            bg = C_ROW_ALT

        mode  = iface["mode"]        or "—"
        desc  = iface["description"] or "—"
        notes = iface["notes"]       or "—"

        if logical:
            ip         = iface["ip_address"].split()[0] if iface["ip_address"] else "—"
            cells_data = [short, desc, mode, ip, notes]
        else:
            acc_nat = iface["access_vlan"] or "—"
            trunk_base = iface["trunk_vlans"] or "—"
            trunk = (trunk_base + f"\nnative: {iface['native_vlan']}"
                     if iface["native_vlan"] else trunk_base)
            cells_data = [short, desc, mode, acc_nat, trunk, notes]

        row = tbl.add_row()
        set_row_min_height(row, 360)

        for ci, (txt, w) in enumerate(zip(cells_data, COL_W_MM)):
            c = row.cells[ci]
            c.width = Mm(w)
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_bg(c, bg)

            align = (WD_ALIGN_PARAGRAPH.CENTER if ci in CENTER
                     else WD_ALIGN_PARAGRAPH.LEFT)
            color = RGB_AMBER if (ci == 2 and iface["shutdown"]) else RGB_DARK_TEXT
            _write_cell(c, txt, bold=(ci == 0), color=color,
                        font_size=10, align=align)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after  = Pt(4)


def _add_vlan_table(doc, vlans: list):
    """Three-column VLAN table: VLAN ID | НАЗВА | СТАТУС (активний / мертвий)."""
    if not vlans:
        return

    # 22 + 130 + 30 = 182 mm
    COL_W_MM = [22, 130, 30]
    COL_HDR  = ["VLAN", "НАЗВА", "СТАТУС"]

    tbl = doc.add_table(rows=1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_fixed_layout(tbl)
    set_table_width(tbl, USABLE_MM)
    set_table_borders(tbl, color_hex=C_BORDER, sz="4")
    set_cell_margins_table(tbl, top=15, left=60, bottom=15, right=40)

    hdr = tbl.rows[0]
    set_repeat_header_row(hdr)
    set_row_min_height(hdr, 400)

    for ci, (h, w) in enumerate(zip(COL_HDR, COL_W_MM)):
        c = hdr.cells[ci]
        c.width = Mm(w)
        set_cell_bg(c, C_ACCENT_BLUE)
        _write_cell(c, h, bold=True, color=RGB_WHITE, font_size=10,
                    align=WD_ALIGN_PARAGRAPH.CENTER)

    for ri, vlan in enumerate(vlans):
        is_used = vlan.get("used", True)
        if is_used:
            bg         = C_WHITE if ri % 2 == 0 else C_ROW_ALT
            text_color = RGB_DARK_TEXT
            status     = "активний"
        else:
            bg         = C_VLAN_DEAD
            text_color = RGB_DEAD_TEXT
            status     = "—"

        row = tbl.add_row()
        set_row_min_height(row, 360)

        c0 = row.cells[0]
        c0.width = Mm(COL_W_MM[0])
        c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_bg(c0, bg)
        _write_cell(c0, str(vlan["id"]), bold=True, color=text_color,
                    font_size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

        c1 = row.cells[1]
        c1.width = Mm(COL_W_MM[1])
        c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_bg(c1, bg)
        _write_cell(c1, vlan["name"] or "—", color=text_color,
                    font_size=10, align=WD_ALIGN_PARAGRAPH.LEFT)

        c2 = row.cells[2]
        c2.width = Mm(COL_W_MM[2])
        c2.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_bg(c2, bg)
        _write_cell(c2, status, color=text_color,
                    font_size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after  = Pt(4)


def build_docx(parsed: dict, user_inputs: dict, output_path: str):
    doc = Document()

    # Calibri as default — full Cyrillic support on any Windows
    doc.styles["Normal"].font.name = "Calibri"

    section = doc.sections[0]
    section.page_width    = Mm(210)
    section.page_height   = Mm(297)
    section.left_margin   = Mm(14)
    section.right_margin  = Mm(14)
    section.top_margin    = Mm(14)
    section.bottom_margin = Mm(14)

    # Remove the default empty paragraph that python-docx inserts
    if doc.paragraphs:
        p = doc.paragraphs[0]._element
        p.getparent().remove(p)

    hostname = parsed.get("hostname") or "Unknown"
    mgmt_ip  = parsed.get("mgmt_ip") or "—"
    model    = user_inputs.get("model")  or "—"
    serial   = user_inputs.get("serial") or "—"
    updated  = date.today().strftime("%d.%m.%Y")

    physical, logical = classify_interfaces(parsed["interfaces"])

    _add_header_block(doc, hostname)
    _add_info_block(doc, model, serial, mgmt_ip, updated)
    _add_section_header(doc, "ФІЗИЧНІ ПОРТИ")
    _add_interface_table(doc, physical)
    _add_section_header(doc, "PORT-CHANNEL ТА VLAN")
    _add_interface_table(doc, logical, logical=True)

    vlans = parsed.get("vlans", [])
    if vlans:
        _add_section_header(doc, "VLAN DEFINITIONS")
        _add_vlan_table(doc, vlans)

    doc.save(output_path)


# ── Audit report ──────────────────────────────────────────────────────────────
C_SEV_BG  = {"high": "FDE2E2", "medium": "FEF3C7", "low": "F1F5F9"}
RGB_SEV   = {"high": RGBColor(0xB9, 0x1C, 0x1C), "medium": RGBColor(0xB4, 0x53, 0x09),
             "low": RGBColor(0x47, 0x55, 0x69)}
RGB_MUTED = RGBColor(0x64, 0x74, 0x8B)


def _write_runs(cell, runs: list, align=WD_ALIGN_PARAGRAPH.LEFT):
    """runs: [(text, bold, size, color, font)]; '\\n' → новий абзац у клітинці."""
    def _fmt(p):
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(0)
        p.alignment = align

    para = cell.paragraphs[0]
    para.clear()
    _fmt(para)
    for text, bold, size, color, font in runs:
        for i, part in enumerate((text or "").split("\n")):
            if i > 0:
                para = cell.add_paragraph()
                _fmt(para)
            if not part:
                continue
            run = para.add_run(part)
            run.bold = bold
            run.font.size = Pt(size)
            if color:
                run.font.color.rgb = color
            _set_run_font(run, font)


def build_audit_docx(report, parsed: dict, user_inputs: dict, output_path: str):
    """DOCX зі звітом аудиту: шапка, зведення, таблиця «рівень · що · де · як виправити»."""
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Mm(14))
    if doc.paragraphs:
        p = doc.paragraphs[0]._element
        p.getparent().remove(p)

    hostname = parsed.get("hostname") or "Unknown"
    _add_header_block(doc, f"{hostname} — аудит конфігурації")
    _add_info_block(doc, user_inputs.get("model") or report.vendor, user_inputs.get("serial") or "—",
                    parsed.get("mgmt_ip") or "—", date.today().strftime("%d.%m.%Y"))

    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(4)
    r = para.add_run(f"{report.vendor}: {summary_text(report)}")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGB_DARK_TEXT
    _set_run_font(r)
    for note in report.skipped:
        np_ = doc.add_paragraph()
        np_.paragraph_format.space_after = Pt(4)
        nr = np_.add_run(note)
        nr.italic = True
        nr.font.size = Pt(9)
        nr.font.color.rgb = RGB_MUTED
        _set_run_font(nr)

    if not report.findings:
        p2 = doc.add_paragraph()
        r2 = p2.add_run("Недолаштованого не знайдено.")
        r2.font.size = Pt(11)
        _set_run_font(r2)
        doc.save(output_path)
        return

    _add_section_header(doc, "ЩО НЕДОЛАШТОВАНО")
    COL_W_MM = [22, 70, 34, 56]
    COL_HDR  = ["РІВЕНЬ", "ЩО НЕДОЛАШТОВАНО", "ДЕ", "ЯК ВИПРАВИТИ"]
    tbl = doc.add_table(rows=1, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_fixed_layout(tbl)
    set_table_width(tbl, USABLE_MM)
    set_table_borders(tbl, color_hex=C_BORDER, sz="4")
    set_cell_margins_table(tbl, top=40, left=70, bottom=40, right=50)

    hdr = tbl.rows[0]
    set_repeat_header_row(hdr)
    set_row_min_height(hdr, 400)
    for ci, (h, w) in enumerate(zip(COL_HDR, COL_W_MM)):
        c = hdr.cells[ci]
        c.width = Mm(w)
        set_cell_bg(c, C_ACCENT_BLUE)
        _write_cell(c, h, bold=True, color=RGB_WHITE, font_size=10,
                    align=WD_ALIGN_PARAGRAPH.CENTER if ci == 0 else WD_ALIGN_PARAGRAPH.LEFT)

    for ri, f in enumerate(report.findings):
        row = tbl.add_row()
        bg = C_WHITE if ri % 2 == 0 else C_ROW_ALT
        cells = row.cells
        for ci, w in enumerate(COL_W_MM):
            cells[ci].width = Mm(w)
            cells[ci].vertical_alignment = WD_ALIGN_VERTICAL.TOP
            set_cell_bg(cells[ci], bg)
        set_cell_bg(cells[0], C_SEV_BG[f.sev])
        cells[0].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _write_cell(cells[0], SEV_LABEL[f.sev], bold=True, color=RGB_SEV[f.sev],
                    font_size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        _write_runs(cells[1], [(f.title, True, 10, RGB_DARK_TEXT, "Calibri"),
                               ("\n" + f.why, False, 9, RGB_MUTED, "Calibri")])
        _write_cell(cells[2], "\n".join(f.where) or "—", color=RGB_DARK_TEXT, font_size=9)
        _write_runs(cells[3], [(f.fix or "—", False, 8, RGB_DARK_TEXT, "Consolas")])

    doc.save(output_path)


# ═════════════════════════════════════════════════════════════════════════════
#  GUI
# ═════════════════════════════════════════════════════════════════════════════

try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
    _BASE = TkinterDnD.Tk
    _HAS_DND = True
except ImportError:
    _BASE = tk.Tk
    _HAS_DND = False
    DND_FILES = None

def _asset(name: str) -> str:
    """Resolve path to an asset file — works both in dev and PyInstaller .exe."""
    base = getattr(sys, '_MEIPASS', os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, 'assets', name)

# Dark theme palette (matches reference style)
_G_BG      = "#282828"
_G_IN      = "#1e1e1e"
_G_BTN     = "#1a1a1a"
_G_FG      = "#e0e0e0"
_G_LBL     = "#b8b8b8"
_G_ACCENT  = "#00ccff"
_G_ACCH    = "#55ddff"
_G_BORDER  = "#3a3a3a"
_G_STATUS  = "#94a3b8"
_G_GREEN   = "#22c55e"
_G_GREENH  = "#4ade80"
_G_GREENBG = "#0f2a1a"


class App(_BASE):
    def __init__(self):
        super().__init__()
        self.title("network-documentr — Cisco / Fortinet DOCX Generator")
        self.resizable(False, True)
        self.geometry("560x155")
        self.configure(bg=_G_BG)
        try:
            self.iconbitmap(_asset("app_icon.ico"))
        except Exception:
            pass  # не критично — іконка просто не відобразиться

        self._paths: list = []
        self._config_entries: list = []
        self._status = tk.StringVar(
            value="Оберіть файл конфігурації та заповніть поля.")

        self._setup_styles()
        self._close_img = self._load_icon("close.png")
        self._build_ui()
        self.update_idletasks()
        self.geometry(f"560x{self.winfo_reqheight()}")
        self.after(50, self._dark_titlebar)
        # Auto-load files passed as CLI arguments (e.g. drag-onto-icon on Windows)
        argv_files = [p for p in sys.argv[1:] if os.path.isfile(p)]
        if argv_files:
            self.after(100, lambda: self._load_files(argv_files))

    def _dark_titlebar(self):
        """Dark title bar on Windows 10 1903+ / Windows 11."""
        try:
            import ctypes
            self.update_idletasks()
            hwnd = ctypes.windll.user32.GetParent(int(self.winfo_id()))
            if not hwnd:
                hwnd = int(self.winfo_id())
            val = ctypes.c_int(1)
            ctypes.windll.dwmapi.DwmSetWindowAttribute(
                hwnd, 20, ctypes.byref(val), ctypes.sizeof(val))
        except Exception:
            pass

    def _load_icon(self, name: str):
        try:
            return tk.PhotoImage(file=_asset(name))
        except Exception:
            return None

    def _setup_styles(self):
        s = ttk.Style(self)
        s.theme_use("clam")

        s.configure(".",
            background=_G_BG, foreground=_G_FG,
            font=("Arial", 10), borderwidth=0, relief="flat")
        s.configure("TFrame",  background=_G_BG)
        s.configure("TLabel",
            background=_G_BG, foreground=_G_LBL,
            font=("Arial", 10, "bold"))

        s.configure("TEntry",
            fieldbackground=_G_IN, foreground=_G_FG,
            bordercolor=_G_BORDER, lightcolor=_G_BORDER, darkcolor=_G_BORDER,
            insertcolor=_G_FG, padding=5)
        s.map("TEntry",
            bordercolor=[("focus", _G_ACCENT)],
            lightcolor=[("focus", _G_ACCENT)])

        s.configure("TButton",
            background=_G_BTN, foreground=_G_ACCENT,
            bordercolor=_G_ACCENT, lightcolor=_G_ACCENT, darkcolor=_G_BTN,
            relief="solid", font=("Arial", 10, "bold"),
            padding=(12, 6), borderwidth=2)
        s.map("TButton",
            background=[("active", "#2a2a2a"), ("pressed", "#111111")],
            foreground=[("active", _G_ACCH)],
            bordercolor=[("active", _G_ACCH)],
            lightcolor=[("active", _G_ACCH)])

        s.configure("TSeparator", background=_G_BORDER)
        s.configure("Status.TLabel",
            background=_G_BG, foreground=_G_STATUS, font=("Arial", 9))
        s.configure("Filename.TLabel",
            background=_G_BG, foreground=_G_FG, font=("Arial", 12))

        s.configure("X.TButton",
            background=_G_BG, foreground="#ff6b6b",
            borderwidth=0, relief="flat", font=("Arial", 20, "bold"),
            padding=(4, 0))
        s.map("X.TButton",
            foreground=[("active", "#ff4444")],
            background=[("active", _G_BG), ("pressed", _G_BG)])

        s.configure("Loaded.TButton",
            background=_G_GREENBG, foreground=_G_GREEN,
            bordercolor=_G_GREEN, lightcolor=_G_GREEN, darkcolor=_G_GREENBG,
            relief="solid", font=("Arial", 10, "bold"),
            padding=(12, 6), borderwidth=2)
        s.map("Loaded.TButton",
            background=[("active", "#1a3a2a"), ("pressed", "#0a1a10")],
            foreground=[("active", _G_GREENH)],
            bordercolor=[("active", _G_GREENH)],
            lightcolor=[("active", _G_GREENH)])

        s.configure("TNotebook", background=_G_BG, borderwidth=0, tabmargins=(0, 0, 0, 0))
        s.configure("TNotebook.Tab",
            background=_G_BTN, foreground=_G_LBL, padding=(12, 5),
            borderwidth=0, font=("Arial", 10, "bold"))
        s.map("TNotebook.Tab",
            background=[("selected", _G_IN)],
            foreground=[("selected", _G_ACCENT)])
        s.configure("Vertical.TScrollbar",
            background=_G_BTN, troughcolor=_G_IN, bordercolor=_G_BG,
            arrowcolor=_G_LBL, lightcolor=_G_BTN, darkcolor=_G_BTN)
        s.configure("Summary.TLabel",
            background=_G_BG, foreground=_G_FG, font=("Arial", 11, "bold"))

        s.configure("Blue.TButton",
            background="#0078d4", foreground="#ffffff",
            bordercolor="#0078d4", lightcolor="#0078d4", darkcolor="#0078d4",
            relief="solid", font=("Arial", 10, "bold"),
            padding=(12, 6), borderwidth=0)
        s.map("Blue.TButton",
            background=[("active", "#106ebe"), ("pressed", "#005a9e")],
            foreground=[("active", "#ffffff")],
            bordercolor=[("active", "#106ebe")],
            lightcolor=[("active", "#106ebe")])

    def _build_ui(self):
        PAD = {"padx": 10, "pady": 5}
        f = ttk.Frame(self, padding="14 12 14 10")
        f.pack(fill="both", expand=True)
        f.columnconfigure(1, weight=1)

        self._f = f
        self._dynamic_widgets: list = []

        self._btn_load = ttk.Button(f, text="Завантажити конфігурації",
                                    style="Blue.TButton", command=self._browse,
                                    cursor="hand2")
        self._btn_load.grid(row=0, column=0, columnspan=3, sticky="ew", **PAD)
        if _HAS_DND:
            self.drop_target_register(DND_FILES)
            self.dnd_bind("<<Drop>>", self._on_drop)

        ttk.Separator(f, orient="horizontal").grid(
            row=999, column=0, columnspan=3, sticky="ew", pady=8)

        ttk.Label(f, textvariable=self._status, style="Status.TLabel").grid(
            row=1000, column=0, columnspan=2, sticky="w", padx=10)
        self._actions = ttk.Frame(f)
        self._btn_audit = ttk.Button(self._actions, text="Аудит",
                                     style="TButton", command=self._audit,
                                     cursor="hand2")
        self._btn_audit.pack(side="left", padx=(0, 8))
        self._btn_generate = ttk.Button(self._actions, text="Створити DOCX",
                                        style="Blue.TButton", command=self._generate,
                                        cursor="hand2")
        self._btn_generate.pack(side="left")
        self._actions.grid(row=1000, column=2, padx=10, pady=2, sticky="e")
        self._actions.grid_remove()

    def _rebuild_entries_ui(self):
        for w in self._dynamic_widgets:
            w.destroy()
        self._dynamic_widgets = []

        PAD = {"padx": 10, "pady": 3}
        f = self._f
        row = 1
        for i, entry in enumerate(self._config_entries):
            if i > 0:
                w = ttk.Separator(f, orient="horizontal")
                w.grid(row=row, column=0, columnspan=3, sticky="ew", pady=4)
                self._dynamic_widgets.append(w)
                row += 1

            filename = Path(entry["path"]).name
            w = ttk.Label(f, text=filename, style="Filename.TLabel")
            w.grid(row=row, column=0, columnspan=2, sticky="w", padx=10, pady=(8, 0))
            self._dynamic_widgets.append(w)
            _xkw = {"image": self._close_img} if self._close_img else {"text": "×"}
            w = ttk.Button(f, **_xkw, style="X.TButton", cursor="hand2",
                           command=lambda idx=i: self._remove_entry(idx))
            w.grid(row=row, column=2, sticky="e", padx=(0, 10), pady=(6, 0))
            self._dynamic_widgets.append(w)
            row += 1

            status_text = format_load_status(entry["parsed"])
            if entry.get("audit"):
                status_text += "  ·  " + short_summary(entry["audit"])
            w = ttk.Label(f, text=status_text, style="Status.TLabel")
            w.grid(row=row, column=0, columnspan=3, sticky="w", padx=10, pady=(0, 6))
            self._dynamic_widgets.append(w)
            row += 1

            w = ttk.Label(f, text="Модель / Вендор:")
            w.grid(row=row, column=0, sticky="w", **PAD)
            self._dynamic_widgets.append(w)
            w = ttk.Entry(f, textvariable=entry["model"], width=40)
            w.grid(row=row, column=1, columnspan=2, sticky="ew", padx=(0, 10), pady=3)
            self._dynamic_widgets.append(w)
            row += 1

            w = ttk.Label(f, text="Серійний номер:")
            w.grid(row=row, column=0, sticky="w", **PAD)
            self._dynamic_widgets.append(w)
            w = ttk.Entry(f, textvariable=entry["serial"], width=40)
            w.grid(row=row, column=1, columnspan=2, sticky="ew", padx=(0, 10), pady=3)
            self._dynamic_widgets.append(w)
            row += 1

        self.update_idletasks()
        self.geometry(f"560x{self.winfo_reqheight()}")

    def _remove_entry(self, index: int):
        self._config_entries.pop(index)
        self._paths.pop(index)
        n = len(self._config_entries)
        if n == 0:
            self._btn_load.configure(text="Завантажити конфігурації", style="Blue.TButton")
            self._actions.grid_remove()
            self._status.set("Оберіть файл конфігурації та заповніть поля.")
            self._rebuild_entries_ui()
            self.update_idletasks()
            self.geometry(f"560x{self.winfo_reqheight()}")
        else:
            self._btn_load.configure(text=f"Завантажено конфігурацій: {n}", style="Loaded.TButton")
            self._rebuild_entries_ui()

    def _browse(self):
        paths = filedialog.askopenfilenames(
            title="Оберіть файл(и) show running-config / show full-configuration",
            filetypes=[("Config files", "*.txt *.conf *.cfg *.log"),
                       ("All files", "*.*")],
        )
        if not paths:
            return
        self._load_files(list(paths))

    def _load_files(self, paths: list):
        try:
            texts = [read_file_safe(p) for p in paths]
            parsed_list = [parse_config(t) for t in texts]
            detected_list = [parse_version_inventory(t) for t in texts]
            audit_list = [run_audit(t, p) for t, p in zip(texts, parsed_list)]
        except Exception:
            self._paths = []
            self._config_entries = []
            self._btn_load.configure(text="Завантажити конфігурації", style="Blue.TButton")
            self._actions.grid_remove()
            self._rebuild_entries_ui()
            self._status.set("Помилка читання файлу.")
            return

        self._paths = list(paths)
        self._config_entries = [
            {
                "path":   p,
                "parsed": parsed,
                "model":  tk.StringVar(value=det["model"]),
                "serial": tk.StringVar(value=det["serial"]),
                "audit":  rep,
            }
            for p, parsed, det, rep in zip(paths, parsed_list, detected_list, audit_list)
        ]
        n = len(paths)
        self._btn_load.configure(text=f"Завантажено конфігурацій: {n}", style="Loaded.TButton")
        self._actions.grid(row=1000, column=2, padx=10, pady=2, sticky="e")
        self._status.set("")
        self._rebuild_entries_ui()

    def _on_drop(self, event):
        tokens = re.findall(r'\{[^}]+\}|\S+', event.data)
        paths = [t.strip('{}') for t in tokens if os.path.isfile(t.strip('{}'))]
        if paths:
            self._load_files(paths)

    def _audit(self):
        if not self._config_entries:
            return
        AuditWindow(self, self._config_entries)

    def _generate(self):
        if not self._config_entries:
            messagebox.showerror("Помилка",
                                 "Оберіть файл(и) конфігурації "
                                 "(show running-config / show full-configuration).")
            return

        if len(self._config_entries) == 1:
            entry = self._config_entries[0]
            src = entry["path"]
            parsed = entry["parsed"]
            user_inputs = {
                "model":  entry["model"].get().strip(),
                "serial": entry["serial"].get().strip(),
            }
            hostname = parsed.get("hostname") or Path(src).stem
            save_path = filedialog.asksaveasfilename(
                title="Зберегти DOCX як…",
                initialdir=str(Path(src).parent),
                initialfile=f"{hostname}_doc.docx",
                defaultextension=".docx",
                filetypes=[("Word document", "*.docx")],
            )
            if not save_path:
                return

            try:
                self._status.set("Генерую документ…")
                self.update()
                build_docx(parsed, user_inputs, save_path)
                physical, logical = classify_interfaces(parsed["interfaces"])
                self._status.set(f"Готово: {Path(save_path).name}")
                messagebox.showinfo(
                    "Готово",
                    f"Документ збережено:\n{save_path}\n\n"
                    f"Hostname       : {hostname}\n"
                    f"Фізичні порти  : {len(physical)}\n"
                    f"Po / VLAN      : {len(logical)}",
                )
            except Exception as exc:
                messagebox.showerror("Помилка генерації", str(exc))
                self._status.set("Помилка. Перевірте файл конфігурації.")

        else:
            out_dir = filedialog.askdirectory(
                title="Оберіть директорію для збереження документів")
            if not out_dir:
                return

            self._status.set("Генерую документи…")
            self.update()
            saved = []
            errors = []
            for entry in self._config_entries:
                try:
                    parsed = entry["parsed"]
                    user_inputs = {
                        "model":  entry["model"].get().strip(),
                        "serial": entry["serial"].get().strip(),
                    }
                    hostname = parsed.get("hostname") or Path(entry["path"]).stem
                    out_path = os.path.join(out_dir, f"{hostname}_doc.docx")
                    build_docx(parsed, user_inputs, out_path)
                    saved.append(Path(out_path).name)
                except Exception as exc:
                    errors.append(f"{Path(entry['path']).name}: {exc}")

            if errors:
                messagebox.showerror("Помилки генерації", "\n".join(errors))
            if saved:
                self._status.set(f"Готово: {len(saved)} файлів")
                messagebox.showinfo(
                    "Готово",
                    f"Збережено {len(saved)} документів:\n" + "\n".join(saved),
                )


_SEV_TAG = {  # tag: (foreground, background)
    "high":   ("#ff8a8a", "#4a1c1c"),
    "medium": ("#fbbf24", "#3d3111"),
    "low":    ("#a8b4c4", "#2a313b"),
}


class AuditWindow(tk.Toplevel):
    """Вікно аудиту: вкладка на кожен завантажений конфіг, знахідки за рівнем."""

    def __init__(self, app, entries: list):
        super().__init__(app)
        self.title("Аудит — що недолаштовано")
        self.configure(bg=_G_BG)
        self.geometry("820x680")
        self.minsize(560, 360)
        try:
            self.iconbitmap(_asset("app_icon.ico"))
        except Exception:
            pass
        self._entries = entries
        self._texts = []

        outer = ttk.Frame(self, padding="12 10 12 10")
        outer.pack(fill="both", expand=True)

        if len(entries) > 1:
            self._nb = ttk.Notebook(outer)
            self._nb.pack(fill="both", expand=True)
            for e in entries:
                tab = ttk.Frame(self._nb, padding="0 8 0 0")
                self._nb.add(tab, text=self._tab_title(e))
                self._build_report(tab, e)
        else:
            self._nb = None
            self._build_report(outer, entries[0])

        bar = ttk.Frame(outer)
        bar.pack(fill="x", pady=(10, 0))
        ttk.Button(bar, text="Копіювати", command=self._copy, cursor="hand2").pack(side="left")
        if len(entries) > 1:
            ttk.Button(bar, text="Зберегти всі…", command=self._save_all,
                       cursor="hand2").pack(side="right", padx=(8, 0))
        ttk.Button(bar, text="Зберегти звіт…", style="Blue.TButton", command=self._save,
                   cursor="hand2").pack(side="right")
        self.after(50, lambda: App._dark_titlebar(self))

    @staticmethod
    def _hostname(e) -> str:
        return e["parsed"].get("hostname") or Path(e["path"]).stem

    def _tab_title(self, e) -> str:
        n = len(e["audit"].findings)
        return f"{self._hostname(e)}  ({n})" if n else f"{self._hostname(e)}  ✓"

    def _build_report(self, parent, e):
        rep = e["audit"]
        ttk.Label(parent, text=f"{self._hostname(e)}  ·  {rep.vendor}  ·  {summary_text(rep)}",
                  style="Summary.TLabel").pack(anchor="w", pady=(0, 8))

        wrap = tk.Frame(parent, bg=_G_BORDER)
        wrap.pack(fill="both", expand=True)
        txt = tk.Text(wrap, wrap="word", bg=_G_IN, fg=_G_FG, relief="flat", borderwidth=0,
                      padx=14, pady=10, font=("Segoe UI", 10), insertbackground=_G_FG,
                      selectbackground="#264f78", cursor="arrow", highlightthickness=0)
        sb = ttk.Scrollbar(wrap, orient="vertical", command=txt.yview)
        txt.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        txt.pack(side="left", fill="both", expand=True, padx=1, pady=1)

        for sev, (fg, bg) in _SEV_TAG.items():
            txt.tag_configure("sev_" + sev, foreground=fg, background=bg, font=("Segoe UI", 9, "bold"))
        txt.tag_configure("title", font=("Segoe UI", 11, "bold"), foreground="#f2f2f2")
        txt.tag_configure("why", foreground="#c4c4c4", lmargin1=18, lmargin2=18)
        txt.tag_configure("lbl", foreground="#7d8796", font=("Segoe UI", 9, "bold"), lmargin1=18)
        txt.tag_configure("where", foreground=_G_FG, font=("Consolas", 10))
        txt.tag_configure("code", font=("Consolas", 10), foreground="#9fe8ff", background="#151a20",
                          lmargin1=30, lmargin2=30, spacing1=0, spacing3=0)
        txt.tag_configure("note", foreground="#fbbf24", font=("Segoe UI", 9, "italic"), spacing3=8)
        txt.tag_configure("ok", foreground=_G_GREEN, font=("Segoe UI", 12, "bold"))
        txt.tag_configure("gap", font=("Segoe UI", 5))

        for note in rep.skipped:
            txt.insert("end", "! " + note + "\n", "note")
        if not rep.findings:
            txt.insert("end", "✓ Недолаштованого не знайдено.\n", "ok")
        for f in rep.findings:
            txt.insert("end", f" {SEV_LABEL[f.sev]} ", "sev_" + f.sev)
            txt.insert("end", "  " + f.title + "\n", "title")
            txt.insert("end", f.why + "\n", "why")
            if f.where:
                txt.insert("end", "Де: ", "lbl")
                txt.insert("end", "; ".join(f.where) + "\n", "where")
            if f.fix:
                txt.insert("end", "Як виправити:\n", "lbl")
                txt.insert("end", f.fix + "\n", "code")
            txt.insert("end", "\n", "gap")
        txt.configure(state="disabled")
        self._texts.append(txt)

    def _current(self):
        if self._nb is None:
            return self._entries[0]
        return self._entries[self._nb.index(self._nb.select())]

    def _copy(self):
        e = self._current()
        self.clipboard_clear()
        self.clipboard_append(report_text(e["audit"], self._hostname(e)))

    @staticmethod
    def _write(e, path: str):
        if path.lower().endswith(".txt"):
            with open(path, "w", encoding="utf-8") as fh:
                fh.write(report_text(e["audit"], AuditWindow._hostname(e)))
        else:
            build_audit_docx(e["audit"], e["parsed"],
                             {"model": e["model"].get().strip(), "serial": e["serial"].get().strip()}, path)

    def _save(self):
        e = self._current()
        path = filedialog.asksaveasfilename(
            parent=self, title="Зберегти звіт аудиту як…",
            initialdir=str(Path(e["path"]).parent),
            initialfile=f"{self._hostname(e)}_audit.docx", defaultextension=".docx",
            filetypes=[("Word document", "*.docx"), ("Текст", "*.txt")])
        if not path:
            return
        try:
            self._write(e, path)
            messagebox.showinfo("Готово", f"Звіт збережено:\n{path}", parent=self)
        except Exception as exc:
            messagebox.showerror("Помилка", str(exc), parent=self)

    def _save_all(self):
        out_dir = filedialog.askdirectory(parent=self, title="Куди зберегти звіти аудиту")
        if not out_dir:
            return
        saved, errors = [], []
        for e in self._entries:
            path = os.path.join(out_dir, f"{self._hostname(e)}_audit.docx")
            try:
                self._write(e, path)
                saved.append(Path(path).name)
            except Exception as exc:
                errors.append(f"{Path(e['path']).name}: {exc}")
        if errors:
            messagebox.showerror("Помилки", "\n".join(errors), parent=self)
        if saved:
            messagebox.showinfo("Готово", f"Збережено {len(saved)}:\n" + "\n".join(saved), parent=self)


if __name__ == "__main__":
    App().mainloop()
